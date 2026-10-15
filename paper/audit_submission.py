import os

try:
    import docx
except ImportError:  # python-docx is optional; only section 4 needs it
    docx = None

# Resolve the submission bundle relative to this file so the audit can be run
# from any working directory.
submission_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'submission')
print(f"=== AUDITING SUBMISSION DIR: {submission_dir} ===")

# 1. Highlights Check
hl_path = os.path.join(submission_dir, 'Highlights.txt')
with open(hl_path, 'r', encoding='utf-8') as f:
    hl_lines = [line.strip() for line in f if line.strip()]

print("\n--- 1. Highlights.txt Check ---")
print(f"Total bullet points: {len(hl_lines)} (Must be exactly 4)")
hl_pass = (len(hl_lines) == 4)
for i, line in enumerate(hl_lines, 1):
    c_len = len(line)
    status = 'PASS' if c_len <= 85 else 'FAIL (>85)'
    if c_len > 85:
        hl_pass = False
    print(f"Line {i}: [{c_len}/85 chars] [{status}] -> {line}")
print(f"Highlights overall status: {'PASSED' if hl_pass else 'FAILED'}")

# 2. Files existence & size check
print("\n--- 2. Required Files Audit ---")
req_files = [
    'Highlights.txt',
    'Cover_Letter.docx',
    'Cover_Letter.txt',
    'Declaration_of_Competing_Interests.txt',
    'CRediT_Author_Statement.txt',
    'manuscript_submission.pdf',
    'manuscript_submission.tex',
    'README_SUBMISSION_GUIDE.txt',
    'fig_tesseract_4d.png',
    'fig_tesseract_architecture.png',
    'fig_key_schedule.png',
    'fig_avalanche_diffusion.png',
    'fig_benchmark_charts.png',
    'fig_related_key_desync.png',
    'Lena_cryptanalysis_result.png',
    'Baboon_cryptanalysis_result.png'
]

missing_count = 0
for rf in req_files:
    fpath = os.path.join(submission_dir, rf)
    if os.path.exists(fpath):
        size = os.path.getsize(fpath)
        print(f"  [FOUND] {rf:<40} ({size:>10,} bytes)")
    else:
        print(f"  [MISSING] {rf}")
        missing_count += 1
print(f"File presence: {'ALL PRESENT' if missing_count == 0 else f'{missing_count} MISSING'}")

# 3. Check Email & Affiliation consistency across text/docx/tex files
print("\n--- 3. Metadata Consistency (Author, Email, Affiliation) ---")
target_email = 'haong.b24at088@stu.ptit.edu.vn'
old_email = 'B24DCAT088@stu.ptit.edu.vn'
target_faculty = 'Faculty of Information Security'
old_faculty = 'Faculty of Information Technology'

text_audit_files = [
    'Cover_Letter.txt',
    'Declaration_of_Competing_Interests.txt',
    'CRediT_Author_Statement.txt',
    'manuscript_submission.tex',
    'README_SUBMISSION_GUIDE.txt'
]

for taf in text_audit_files:
    fpath = os.path.join(submission_dir, taf)
    with open(fpath, 'r', encoding='utf-8', errors='ignore') as f:
        content = f.read()
    has_target_email = target_email in content
    has_old_email = old_email in content
    has_target_fac = target_faculty in content
    has_old_fac = old_faculty in content
    print(f"File: {taf}")
    print(f"  Target email present: {has_target_email} | Old email present: {has_old_email}")
    print(f"  Target faculty present: {has_target_fac} | Old faculty present: {has_old_fac}")

# 4. Check Cover_Letter.docx specifically
print("\n--- 4. Cover_Letter.docx In-depth Audit ---")
if docx is None:
    print("  SKIPPED: python-docx not installed (pip install python-docx)")
    docx_has_email = docx_has_old_email = docx_has_fac = docx_has_old_fac = None
    has_s1 = has_s2 = has_s3 = None
    full_docx_text = ''
else:
    doc = docx.Document(os.path.join(submission_dir, 'Cover_Letter.docx'))
    full_docx_text = '\n'.join([p.text for p in doc.paragraphs])
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                full_docx_text += '\n' + c.text

    docx_has_email = target_email in full_docx_text
    docx_has_old_email = old_email in full_docx_text
    docx_has_fac = target_faculty in full_docx_text
    docx_has_old_fac = old_faculty in full_docx_text

    # Check the 3 key sentence structures in Cover_Letter.docx
    sent1_snippet = "Although conventional literature demonstrates that software AES implementations achieve acceptable throughput"
    sent2_snippet = "To resolve this fundamental operational trade-off, this manuscript introduces Tesseract-128"
    sent3_snippet = "Within a word-level MILP model of the ARX carry chain, automated searches establish lower bounds of 78 active S-boxes"

    has_s1 = sent1_snippet in full_docx_text
    has_s2 = sent2_snippet in full_docx_text
    has_s3 = sent3_snippet in full_docx_text

    print(f"Paragraphs count: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")
    print(f"  Target email: {docx_has_email} | Old email: {docx_has_old_email}")
    print(f"  Target faculty: {docx_has_fac} | Old faculty: {docx_has_old_fac}")
    print(f"  Sentence Structure 1 present: {has_s1}")
print(f"  Sentence Structure 2 present: {has_s2}")
print(f"  Sentence Structure 3 present: {has_s3}")

# 5. Check PDF validity
print("\n--- 5. PDF Check ---")
pdf_path = os.path.join(submission_dir, 'manuscript_submission.pdf')
print(f"PDF exists: {os.path.exists(pdf_path)}")
print(f"PDF size: {os.path.getsize(pdf_path):,} bytes")
try:
    import pypdf
    reader = pypdf.PdfReader(pdf_path)
    print(f"PDF total pages: {len(reader.pages)}")
    first_page_text = reader.pages[0].extract_text()
    print(f"First page contains target faculty: {'Faculty of Information Security' in first_page_text}")
    print(f"First page contains target email: {'haong.b24at088@stu.ptit.edu.vn' in first_page_text}")
except ImportError:
    print("pypdf not installed, checking via file size.")
